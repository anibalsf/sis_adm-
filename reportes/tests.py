"""Tests para asistencia y reportes financieros."""
import pytest
from datetime import date

from tesoreria.models import Pago, Egreso, TipoPago
from afiliados.models import Afiliado


@pytest.fixture
def tipo_ingreso(db):
    return TipoPago.objects.create(
        nombre='Cuota Mensual Test', descripcion='Cuota', tipo='ingreso'
    )


@pytest.fixture
def tipo_egreso(db):
    return TipoPago.objects.create(
        nombre='Mantenimiento Test', descripcion='Mant.', tipo='egreso'
    )


@pytest.fixture
def afiliado(db):
    return Afiliado.objects.create(
        nombres='Juan Perez',
        apellidos='Prueba Uno',
        ci='12345678',
        fecha_ingreso=date(2026, 1, 1),
        estado='activo',
    )


@pytest.mark.django_db
def test_resumen_periodo_solo_transacciones_validas(afiliado, tipo_ingreso, tipo_egreso):
    from reportes.query_helpers import resumen_periodo

    Pago.objects.create(
        afiliado=afiliado, tipo_pago=tipo_ingreso,
        monto=100, fecha_pago=date(2026, 9, 1), estado='completado',
    )
    Pago.objects.create(
        afiliado=afiliado, tipo_pago=tipo_ingreso,
        monto=500, fecha_pago=date(2026, 9, 2), estado='anulado',
        motivo_anulacion='Error',
    )
    Egreso.objects.create(
        fecha=date(2026, 9, 3), monto=50, descripcion='Gasto valido',
        tipo_pago=tipo_egreso, estado='aprobado',
    )
    Egreso.objects.create(
        fecha=date(2026, 9, 4), monto=999, descripcion='Gasto anulado',
        tipo_pago=tipo_egreso, estado='anulado',
    )

    resumen = resumen_periodo()

    # Solo el pago completado y el egreso aprobado cuentan
    assert resumen['total_ingresos'] == 100.0
    assert resumen['total_egresos'] == 50.0
    assert resumen['saldo'] == 50.0
    assert resumen['count_ingresos'] == 1
    assert resumen['count_egresos'] == 1

    # Los anulados quedan explícitamente excluidos y reportados
    excluidos = resumen['anulados_cancelados']
    assert excluidos['count_ingresos'] == 1
    assert excluidos['monto_ingresos'] == 500.0
    assert excluidos['count_egresos'] == 1
    assert excluidos['monto_egresos'] == 999.0


@pytest.mark.django_db
def test_pagos_validos_ignora_cancelados(afiliado, tipo_ingreso):
    from reportes.query_helpers import pagos_validos, pagos_excluidos

    Pago.objects.create(
        afiliado=afiliado, tipo_pago=tipo_ingreso,
        monto=10, fecha_pago=date(2026, 9, 1), estado='completado',
    )
    Pago.objects.create(
        afiliado=afiliado, tipo_pago=tipo_ingreso,
        monto=20, fecha_pago=date(2026, 9, 2), estado='cancelado',
    )

    assert pagos_validos().count() == 1
    assert pagos_excluidos().count() == 1


@pytest.mark.django_db
def test_balance_endpoint_solo_transacciones_validas(authenticated_client, afiliado, tipo_ingreso, tipo_egreso):
    Pago.objects.create(
        afiliado=afiliado, tipo_pago=tipo_ingreso,
        monto=100, fecha_pago=date(2026, 9, 1), estado='completado',
    )
    Pago.objects.create(
        afiliado=afiliado, tipo_pago=tipo_ingreso,
        monto=700, fecha_pago=date(2026, 9, 2), estado='anulado',
    )
    Egreso.objects.create(
        fecha=date(2026, 9, 3), monto=40, descripcion='Gasto',
        tipo_pago=tipo_egreso, estado='aprobado',
    )

    response = authenticated_client.get('/api/reportes/balance')
    assert response.status_code == 200

    data = response.data
    assert data['total_ingresos'] == 100.0
    assert data['total_egresos'] == 40.0
    assert data['saldo'] == 60.0
    assert data['excluidos_por_estado']['count_ingresos'] == 1
    assert data['excluidos_por_estado']['monto_ingresos'] == 700.0


@pytest.mark.django_db
def test_balance_expone_totales_de_todos_los_egresos(authenticated_client, tipo_egreso):
    """El balance debe mostrar el total de TODOS los egresos y su desglose."""
    Egreso.objects.create(
        fecha=date(2026, 9, 1), monto=100, descripcion='Aprobado',
        tipo_pago=tipo_egreso, estado='aprobado',
    )
    Egreso.objects.create(
        fecha=date(2026, 9, 2), monto=250, descripcion='Pendiente',
        tipo_pago=tipo_egreso, estado='pendiente_aprobacion',
    )
    Egreso.objects.create(
        fecha=date(2026, 9, 3), monto=50, descripcion='Anulado',
        tipo_pago=tipo_egreso, estado='anulado',
    )

    response = authenticated_client.get('/api/reportes/balance')

    assert response.status_code == 200
    data = response.data

    # El saldo del balance solo considera los aprobados.
    assert data['total_egresos'] == 100.0
    assert data['count_egresos'] == 1

    totales = data['totales_egresos']
    assert totales['total_monto'] == 400.0
    assert totales['total_aprobado'] == 100.0
    assert totales['total_pendiente_aprobacion'] == 250.0
    assert totales['total_anulado'] == 50.0
    assert totales['count_todos'] == 3
    assert totales['count_aprobado'] == 1
    assert totales['count_pendiente_aprobacion'] == 1
    assert totales['count_anulado'] == 1

    # El listado completo incluye los 3 y marca cuáles son válidos.
    assert len(data['lista_egresos_todos']) == 3
    assert len(data['lista_egresos']) == 1
    validos = [e for e in data['lista_egresos_todos'] if e['es_valido']]
    assert len(validos) == 1


@pytest.mark.django_db
def test_balance_respeta_rango_de_fechas_en_totales(authenticated_client, tipo_egreso):
    Egreso.objects.create(
        fecha=date(2026, 1, 10), monto=999, descripcion='Fuera de rango',
        tipo_pago=tipo_egreso, estado='aprobado',
    )
    Egreso.objects.create(
        fecha=date(2026, 9, 10), monto=100, descripcion='Dentro de rango',
        tipo_pago=tipo_egreso, estado='aprobado',
    )

    response = authenticated_client.get(
        '/api/reportes/balance?fecha_inicio=2026-09-01&fecha_fin=2026-09-30'
    )

    assert response.status_code == 200
    assert response.data['totales_egresos']['total_monto'] == 100.0
    assert response.data['totales_egresos']['count_todos'] == 1


@pytest.mark.django_db
def test_transacciones_endpoint_resumen(authenticated_client, afiliado, tipo_ingreso):
    Pago.objects.create(
        afiliado=afiliado, tipo_pago=tipo_ingreso,
        monto=25, fecha_pago=date(2026, 9, 1), estado='completado',
    )
    Pago.objects.create(
        afiliado=afiliado, tipo_pago=tipo_ingreso,
        monto=60, fecha_pago=date(2026, 9, 2), estado='cancelado',
    )

    response = authenticated_client.get('/api/reportes/transacciones')

    assert response.status_code == 200

    data = response.data
    assert data['count'] == 1
    assert data['resumen']['total_ingresos'] == 25.0
    assert data['resumen']['count_ingresos'] == 1
    assert data['excluidos_por_estado']['ingresos']['count'] == 1


# --------------------------------------------------------------------------- #
# Reporte por categoría (ingresos): cálculo de deudas
# --------------------------------------------------------------------------- #

@pytest.fixture
def otro_afiliado(db):
    return Afiliado.objects.create(
        nombres='Maria', apellidos='Sanchez Rios', ci='87654321',
        telefono='70000001', fecha_ingreso=date(2026, 1, 1), estado='activo',
    )


@pytest.fixture
def afiliado_pasivo(db):
    return Afiliado.objects.create(
        nombres='Pedro', apellidos='Inactivo', ci='11223344',
        fecha_ingreso=date(2026, 1, 1), estado='pasivo',
    )


def _pago(afiliado, tipo_pago, monto, estado='completado', fecha=date(2026, 9, 1)):
    return Pago.objects.create(
        afiliado=afiliado, tipo_pago=tipo_pago,
        monto=monto, fecha_pago=fecha, estado=estado,
    )


@pytest.mark.django_db
def test_pendientes_sin_monto_esperado_solo_lista_quien_no_pago(
    afiliado, otro_afiliado, afiliado_pasivo, tipo_ingreso
):
    from reportes.views import _pendientes_por_categoria

    _pago(otro_afiliado, tipo_ingreso, 100)
    # Un pago en estado no válido no cuenta como cancelación.
    _pago(afiliado, tipo_ingreso, 50, estado='anulado')

    activos = Afiliado.objects.filter(estado='activo', is_active=True)
    resultado = _pendientes_por_categoria(
        Pago, tipo_ingreso, None, None, None, list(activos)
    )

    assert [p['nombre'] for p in resultado['pendientes']] == [afiliado.nombre_completo]
    assert resultado['count_pendientes'] == 1
    assert resultado['count_sin_pago'] == 1
    assert resultado['count_parciales'] == 0
    assert resultado['monto_total_pendiente'] is None
    # El pasivo nunca entra en la lista de deudores.
    assert afiliado_pasivo.id not in [p['afiliado_id'] for p in resultado['pendientes']]


@pytest.mark.django_db
def test_pendientes_con_monto_esperado_detecta_pagos_parciales(
    afiliado, otro_afiliado, tipo_ingreso
):
    from reportes.views import _pendientes_por_categoria

    _pago(afiliado, tipo_ingreso, 30)          # pago parcial: debe 50
    _pago(otro_afiliado, tipo_ingreso, 100)    # pagado por completo: no es deudor

    resultado = _pendientes_por_categoria(
        Pago, tipo_ingreso, None, None, 50,
        list(Afiliado.objects.filter(estado='activo', is_active=True))
    )

    pendientes = {p['afiliado_id']: p for p in resultado['pendientes']}
    assert set(pendientes) == {afiliado.id}
    assert pendientes[afiliado.id]['deuda'] == 20.0
    assert pendientes[afiliado.id]['pagado'] == 30.0
    assert resultado['count_parciales'] == 1
    assert resultado['monto_total_pendiente'] == 20.0


@pytest.mark.django_db
def test_pendientes_sin_monto_esperado_no_inventa_deuda(
    afiliado, otro_afiliado, tipo_ingreso
):
    """Sin monto esperado la deuda es None, pero se informa lo que pagó cada uno."""
    from reportes.views import _pendientes_por_categoria

    _pago(otro_afiliado, tipo_ingreso, 40)  # parcial, pero sin monto esperado

    resultado = _pendientes_por_categoria(
        Pago, tipo_ingreso, None, None, None,
        list(Afiliado.objects.filter(estado='activo', is_active=True))
    )

    pendientes = {p['afiliado_id']: p for p in resultado['pendientes']}
    assert set(pendientes) == {afiliado.id}
    assert pendientes[afiliado.id]['deuda'] is None
    assert pendientes[afiliado.id]['pagado'] == 0.0
    assert resultado['monto_total_pendiente'] is None


@pytest.mark.django_db
def test_reporte_categoria_cuadra_monto_faltante_con_detalle(
    authenticated_client, afiliado, otro_afiliado, tipo_ingreso
):
    # Juan no pagó nada (deuda 50) y Maria pagó 30 de 50 (deuda 20).
    _pago(otro_afiliado, tipo_ingreso, 30)

    response = authenticated_client.get(
        f'/api/reportes/por-categoria/?tipo_pago_id={tipo_ingreso.id}&monto_esperado=50'
    )

    assert response.status_code == 200
    resumen = response.data['resumen']
    pendientes = response.data['pendientes']

    # El KPI de faltante debe ser exactamente la suma de las deudas del detalle.
    assert resumen['monto_faltante'] == sum(p['deuda'] for p in pendientes)
    assert resumen['monto_faltante'] == 70.0
    assert resumen['total_esperado'] == 100.0
    assert resumen['total_recaudado'] == 30.0
    assert resumen['count_pendientes'] == len(pendientes) == 2
    assert resumen['count_parciales'] == 1
    assert resumen['count_sin_pago'] == 1


@pytest.mark.django_db
def test_reporte_categoria_todas_incluye_pendientes_por_categoria(
    authenticated_client, afiliado, otro_afiliado, tipo_ingreso
):
    _pago(otro_afiliado, tipo_ingreso, 50)

    response = authenticated_client.get(
        '/api/reportes/por-categoria/?todas=1&monto_esperado=50'
    )

    assert response.status_code == 200
    categorias = {c['categoria']['id']: c for c in response.data['categorias']}
    categoria = categorias[tipo_ingreso.id]

    assert [p['nombre'] for p in categoria['pendientes']] == [afiliado.nombre_completo]
    assert categoria['monto_faltante'] == 50.0
    assert categoria['count_sin_pago'] == 1
    assert response.data['resumen']['total_pendientes'] == 1


@pytest.mark.django_db
def test_reporte_categoria_todas_total_faltante_cuadra_con_categorias(
    authenticated_client, afiliado, otro_afiliado, afiliado_pasivo, tipo_ingreso
):
    # El pasivo no cuenta para el total esperado, pero su pago sí suma como
    # recaudado: por eso el faltante debe venir de la deuda real por afiliado.
    _pago(afiliado_pasivo, tipo_ingreso, 100)

    response = authenticated_client.get('/api/reportes/por-categoria/?todas=1&monto_esperado=50')

    assert response.status_code == 200
    categorias = response.data['categorias']
    assert response.data['resumen']['monto_faltante'] == sum(
        c['monto_faltante'] for c in categorias
    )
    # 2 activos sin pagar = 100 de deuda real (no 0 por el pago del pasivo).
    assert response.data['resumen']['monto_faltante'] == 100.0
    assert response.data['resumen']['total_pendientes'] == 2


@pytest.mark.django_db
def test_reporte_categoria_detecta_pagos_duplicados_con_detalle(
    authenticated_client, afiliado, tipo_ingreso
):
    _pago(afiliado, tipo_ingreso, 50, fecha=date(2026, 9, 1))
    _pago(afiliado, tipo_ingreso, 50, fecha=date(2026, 9, 5))

    response = authenticated_client.get(
        f'/api/reportes/por-categoria/?tipo_pago_id={tipo_ingreso.id}'
    )

    assert response.status_code == 200
    assert response.data['resumen']['count_afiliados_con_pago_duplicado'] == 1

    duplicados = response.data['duplicados']
    assert len(duplicados) == 1
    assert duplicados[0]['nombre'] == afiliado.nombre_completo
    assert duplicados[0]['cantidad_pagos'] == 2
    assert duplicados[0]['total_pagado'] == 100.0
    # El detalle incluye los pagos individuales a revisar.
    assert len(duplicados[0]['pagos']) == 2
    assert [p['monto'] for p in duplicados[0]['pagos']] == [50.0, 50.0]


@pytest.mark.django_db
def test_reporte_categoria_todas_agrega_duplicados_globales(
    authenticated_client, afiliado, tipo_ingreso
):
    _pago(afiliado, tipo_ingreso, 50)
    _pago(afiliado, tipo_ingreso, 30)

    response = authenticated_client.get('/api/reportes/por-categoria/?todas=1')

    assert response.status_code == 200
    assert len(response.data['duplicados']) == 1
    duplicado = response.data['duplicados'][0]
    assert duplicado['categoria']['id'] == tipo_ingreso.id
    assert duplicado['total_pagado'] == 80.0
    assert response.data['resumen']['total_monto_duplicado'] == 80.0
    assert response.data['categorias'][0]['duplicados'][0]['cantidad_pagos'] == 2


@pytest.mark.django_db
def test_reporte_categoria_duplicados_ignora_pagos_no_completados(
    authenticated_client, afiliado, tipo_ingreso
):
    _pago(afiliado, tipo_ingreso, 50)
    _pago(afiliado, tipo_ingreso, 50, estado='anulado')

    response = authenticated_client.get(
        f'/api/reportes/por-categoria/?tipo_pago_id={tipo_ingreso.id}'
    )

    assert response.status_code == 200
    assert response.data['duplicados'] == []
    assert response.data['resumen']['count_afiliados_con_pago_duplicado'] == 0


@pytest.mark.django_db
def test_reporte_categoria_lista_categorias_sin_filtro(authenticated_client, tipo_ingreso):
    response = authenticated_client.get('/api/reportes/por-categoria/')

    assert response.status_code == 200
    assert [t['nombre'] for t in response.data['tipos_pago']] == [tipo_ingreso.nombre]


@pytest.mark.django_db
def test_excel_categoria_incluye_columna_deuda(
    authenticated_client, afiliado, otro_afiliado, tipo_ingreso
):
    import io

    from openpyxl import load_workbook

    _pago(otro_afiliado, tipo_ingreso, 20)

    response = authenticated_client.get(
        f'/api/reportes/por-categoria/excel/?tipo_pago_id={tipo_ingreso.id}&monto_esperado=50'
    )

    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb['Pendientes']
    encabezados = [celda.value for celda in ws[1]]
    assert 'Deuda (Bs)' in encabezados
    assert ws.cell(row=2, column=encabezados.index('Nombre Completo') + 1).value == afiliado.nombre_completo


@pytest.mark.django_db
def test_pdf_categoria_se_genera_con_deudas(authenticated_client, afiliado, otro_afiliado, tipo_ingreso):
    _pago(otro_afiliado, tipo_ingreso, 20)

    response = authenticated_client.get(
        f'/api/reportes/por-categoria/pdf/?tipo_pago_id={tipo_ingreso.id}&monto_esperado=50'
    )

    assert response.status_code == 200
    assert response['Content-Type'] == 'application/pdf'
    assert response.content.startswith(b'%PDF')


@pytest.mark.django_db
def test_reporte_egresos_por_categoria(authenticated_client, tipo_egreso):
    Egreso.objects.create(
        fecha=date(2026, 9, 5), monto=300, descripcion='Reparacion',
        tipo_pago=tipo_egreso, estado='aprobado',
    )

    response = authenticated_client.get(
        f'/api/reportes/por-categoria/?tipo=egreso&tipo_pago_id={tipo_egreso.id}'
    )

    assert response.status_code == 200
    assert response.data['tipo'] == 'egreso'
    assert response.data['resumen']['total_egresos'] == 300.0
    assert response.data['resumen']['count_egresos'] == 1
    assert len(response.data['egresos']) == 1