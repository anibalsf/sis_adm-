from rest_framework import viewsets, filters
from rest_framework.permissions import BasePermission, SAFE_METHODS
from rest_framework.decorators import action
from rest_framework.response import Response
from django.conf import settings
from django.http import HttpResponse
from django.db.models import Count, Q
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .models import Asistencia
from .serializers import AsistenciaSerializer
from sanciones.models import Sancion


class AsistenciaViewSet(viewsets.ModelViewSet):
    queryset = Asistencia.objects.select_related('reunion', 'afiliado').all()
    serializer_class = AsistenciaSerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['reunion__tema', 'afiliado__ci']
    ordering_fields = ['created_at']

    class IsSecretariaOrDirectivaOrReadOnly(BasePermission):
        def has_permission(self, request, view):
            if request.method in SAFE_METHODS:
                return True
            user = request.user
            if not user or not user.is_authenticated:
                return False
            if user.is_superuser or user.is_staff:
                return True
            return user.groups.filter(name__in=['Secretaria', 'Directiva']).exists()

    permission_classes = [IsSecretariaOrDirectivaOrReadOnly]

    def get_queryset(self):
        """Filtrar asistencias por reunion_id si se proporciona"""
        queryset = super().get_queryset()
        reunion_id = self.request.query_params.get('reunion')
        if reunion_id:
            queryset = queryset.filter(reunion_id=reunion_id)
        return queryset

    def perform_create(self, serializer):
        """Crear asistencia y generar sanción automática si hay falta"""
        instance = serializer.save()
        
        # Sincronizar presente boolean (deprecated)
        if instance.estado == 'presente':
            instance.presente = True
        else:
            instance.presente = False
        instance.save(update_fields=['presente'])

        if instance.estado == 'falta':
            # Crear sanción por falta
            Sancion.objects.get_or_create(
                afiliado=instance.afiliado,
                tipo='falta_reunion',
                referencia_asistencia=instance,
                defaults={
                    'motivo': f'Inasistencia a reunión: {instance.reunion.tema}',
                    'monto': getattr(settings, 'SANCTION_ABSENCE_AMOUNT', 50),
                    'estado': 'pendiente',
                },
            )

    def perform_update(self, serializer):
        """Actualizar asistencia y manejar sanciones según el cambio"""
        old_instance = self.get_object()
        was_falta = old_instance.estado == 'falta'
        
        instance = serializer.save()
        is_falta = instance.estado == 'falta'
        
        # Sincronizar presente boolean (deprecated)
        if instance.estado == 'presente':
            instance.presente = True
        else:
            instance.presente = False
        instance.save(update_fields=['presente'])
        
        # Si cambió a falta, crear sanción
        if not was_falta and is_falta:
            Sancion.objects.get_or_create(
                afiliado=instance.afiliado,
                tipo='falta_reunion',
                referencia_asistencia=instance,
                defaults={
                    'motivo': f'Inasistencia a reunión: {instance.reunion.tema}',
                    'monto': getattr(settings, 'SANCTION_ABSENCE_AMOUNT', 50),
                    'estado': 'pendiente',
                },
            )
        # Si dejó de ser falta (pasó a presente, permiso o atraso), anular sanción pendiente
        elif was_falta and not is_falta:
            Sancion.objects.filter(
                referencia_asistencia=instance,
                estado='pendiente'
            ).update(estado='anulada')
    
    @action(detail=False, methods=['get'])
    def exportar_excel(self, request):
        """
        Exporta asistencias a Excel con filtros opcionales
        GET /api/asistencias/exportar_excel/?fecha_inicio=YYYY-MM-DD&fecha_fin=YYYY-MM-DD&afiliado=ID
        """
        # Obtener parámetros
        fecha_inicio = request.query_params.get('fecha_inicio')
        fecha_fin = request.query_params.get('fecha_fin')
        afiliado_id = request.query_params.get('afiliado')
        
        # Filtrar asistencias
        queryset = self.get_queryset()
        
        if fecha_inicio:
            queryset = queryset.filter(reunion__fecha__gte=fecha_inicio)
        if fecha_fin:
            queryset = queryset.filter(reunion__fecha__lte=fecha_fin)
        if afiliado_id:
            queryset = queryset.filter(afiliado_id=afiliado_id)
        
        queryset = queryset.select_related('afiliado', 'reunion').order_by('-reunion__fecha')
        
        # Crear workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Asistencias"
        
        # Estilos
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
        center_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Encabezados
        headers = ['ID', 'Afiliado', 'CI', 'Reunión', 'Fecha', 'Estado', 'Justificación']
        ws.append(headers)
        
        # Aplicar estilo a encabezados
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
            cell.border = border
        
        # Datos
        for asistencia in queryset:
            ws.append([
                asistencia.id,
                asistencia.afiliado.nombre_completo if asistencia.afiliado else 'N/A',
                asistencia.afiliado.ci if asistencia.afiliado else '',
                asistencia.reunion.tipo if asistencia.reunion else 'N/A',
                asistencia.reunion.fecha.strftime('%d/%m/%Y') if asistencia.reunion else '',
                'PRESENTE' if asistencia.presente else 'AUSENTE',
                asistencia.observaciones or ''
            ])
        
        # Ajustar ancho de columnas
        column_widths = [8, 30, 12, 20, 12, 12, 40]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + i)].width = width
        
        # Aplicar bordes y colores a todas las celdas de datos
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = border
                if cell.column == 6:  # Columna Estado
                    if cell.value == 'PRESENTE':
                        cell.fill = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
                    else:
                        cell.fill = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
        
        # Preparar respuesta
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"asistencias_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        wb.save(response)
        return response
    
    @action(detail=False, methods=['get'])
    def inasistencias_frecuentes(self, request):
        """
        Retorna afiliados con múltiples inasistencias
        GET /api/asistencias/inasistencias_frecuentes/?periodo_meses=6&minimo_faltas=3
        """
        # Parámetros
        periodo_meses = int(request.query_params.get('periodo_meses', 6))
        minimo_faltas = int(request.query_params.get('minimo_faltas', 3))
        
        fecha_desde = datetime.now().date() - timedelta(days=periodo_meses * 30)
        
        # Agrupar por afiliado
        from afiliados.models import Afiliado
        afiliados = Afiliado.objects.annotate(
            total_asistencias=Count('asistencia', filter=Q(asistencia__reunion__fecha__gte=fecha_desde)),
            total_faltas=Count('asistencia', filter=Q(
                asistencia__reunion__fecha__gte=fecha_desde,
                asistencia__presente=False
            )),
            faltas_injustificadas=Count('asistencia', filter=Q(
                asistencia__reunion__fecha__gte=fecha_desde,
                asistencia__presente=False,
                asistencia__observaciones__isnull=True
            ))
        ).filter(
            total_faltas__gte=minimo_faltas
        ).order_by('-total_faltas')
        
        # Serializar resultados
        resultados = []
        for afiliado in afiliados:
            porcentaje_asistencia = 0
            if afiliado.total_asistencias > 0:
                porcentaje_asistencia = ((afiliado.total_asistencias - afiliado.total_faltas) / afiliado.total_asistencias) * 100
            
            resultados.append({
                'afiliado_id': afiliado.id,
                'nombre_completo': afiliado.nombre_completo,
                'ci': afiliado.ci,
                'total_asistencias': afiliado.total_asistencias,
                'total_faltas': afiliado.total_faltas,
                'faltas_injustificadas': afiliado.faltas_injustificadas,
                'porcentaje_asistencia': round(porcentaje_asistencia, 2)
            })
        
        return Response({
            'periodo_meses': periodo_meses,
            'fecha_desde': fecha_desde,
            'cantidad_afiliados': len(resultados),
            'afiliados': resultados
        })
    
    @action(detail=False, methods=['get'])
    def estadisticas(self, request):
        """
        Estadísticas generales de asistencia
        GET /api/asistencias/estadisticas/?fecha_inicio=YYYY-MM-DD&fecha_fin=YYYY-MM-DD
        """
        # Parámetros de filtro
        fecha_inicio = request.query_params.get('fecha_inicio')
        fecha_fin = request.query_params.get('fecha_fin')
        
        queryset = self.get_queryset()
        
        if fecha_inicio:
            queryset = queryset.filter(reunion__fecha__gte=fecha_inicio)
        if fecha_fin:
            queryset = queryset.filter(reunion__fecha__lte=fecha_fin)
        
        # Estadísticas globales
        total = queryset.count()
        presentes = queryset.filter(presente=True).count()
        ausentes = queryset.filter(presente=False).count()
        porcentaje = (presentes / total * 100) if total > 0 else 0
        
        # Por afiliado
        from afiliados.models import Afiliado
        por_afiliado = []
        
        afiliados_stats = Afiliado.objects.annotate(
            total_asist=Count('asistencia', filter=Q(asistencia__in=queryset)),
            total_pres=Count('asistencia', filter=Q(asistencia__in=queryset, asistencia__presente=True)),
            total_aus=Count('asistencia', filter=Q(asistencia__in=queryset, asistencia__presente=False))
        ).filter(total_asist__gt=0)
        
        for afiliado in afiliados_stats:
            porc = (afiliado.total_pres / afiliado.total_asist * 100) if afiliado.total_asist > 0 else 0
            por_afiliado.append({
                'afiliado_id': afiliado.id,
                'nombre': afiliado.nombre_completo,
                'total': afiliado.total_asist,
                'presentes': afiliado.total_pres,
                'ausentes': afiliado.total_aus,
                'porcentaje': round(porc, 2)
            })
        
        return Response({
            'total_asistencias': total,
            'presentes': presentes,
            'ausentes': ausentes,
            'porcentaje_asistencia': round(porcentaje, 2),
            'por_afiliado': sorted(por_afiliado, key=lambda x: x['porcentaje'])
        })