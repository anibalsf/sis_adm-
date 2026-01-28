"""
Fase 5: Asistencia Mejorada
Código para mejorar el módulo de asistencias
"""

# ========================================
# INSTALACIÓN REQUERIDA
# ========================================
"""
pip install openpyxl
"""

# ========================================
# IMPORTS NECESARIOS
# ========================================
"""
from rest_framework.decorators import action
from rest_framework.response import Response
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, timedelta
from django.db.models import Count, Q
"""

# ========================================
# MÉTODO 1: Exportar asistencias a Excel
# ========================================
"""
@action(detail=False, methods=['get'])
def exportar_excel(self, request):
    '''
    Exportar asistencias a Excel con formato profesional.
    URL: GET /api/asistencias/exportar_excel/?fecha_desde=2025-01-01&fecha_hasta=2025-12-31
    '''
    from asistencias.models import Asistencia
    
    # Obtener parámetros
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')
    reunion_id = request.GET.get('reunion_id')
    
    # Filtrar asistencias
    asistencias = Asistencia.objects.select_related('afiliado', 'reunion').all()
    
    if fecha_desde:
        asistencias = asistencias.filter(reunion__fecha__gte=fecha_desde)
    if fecha_hasta:
        asistencias = asistencias.filter(reunion__fecha__lte=fecha_hasta)
    if reunion_id:
        asistencias = asistencias.filter(reunion_id=reunion_id)
    
    # Crear workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencias"
    
    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Encabezados
    headers = ['Nº', 'Reunión', 'Fecha', 'Afiliado', 'CI', 'Presente', 'Observaciones']
    ws.append(headers)
    
    # Aplicar estilos a encabezados
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    # Datos
    for idx, asistencia in enumerate(asistencias, start=2):
        reunion_nombre = asistencia.reunion.tema if hasattr(asistencia.reunion, 'tema') else f"Reunión {asistencia.reunion.id}"
        fecha = asistencia.reunion.fecha if hasattr(asistencia.reunion, 'fecha') else 'N/A'
        afiliado = asistencia.afiliado.nombre_completo if asistencia.afiliado else 'N/A'
        ci = asistencia.afiliado.ci if asistencia.afiliado else 'N/A'
        presente = 'SÍ' if asistencia.presente else 'NO'
        observaciones = getattr(asistencia, 'observaciones', '') or ''
        
        row = [idx-1, reunion_nombre, str(fecha), afiliado, ci, presente, observaciones]
        ws.append(row)
        
        # Aplicar bordes
        for col_num in range(1, len(headers) + 1):
            ws.cell(row=idx, column=col_num).border = border
        
        # Colorear fila si no asistió
        if not asistencia.presente:
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=idx, column=col_num)
                cell.fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", filltype="solid")
    
    # Ajustar ancho de columnas
    column_widths = [5, 30, 12, 30, 12, 10, 40]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    
    # Guardar en buffer
    from io import BytesIO
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    # Crear respuesta HTTP
    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'asistencias_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename={filename}'
    
    return response
"""

# ========================================
# MÉTODO 2: Reporte de inasistencias frecuentes
# ========================================
"""
@action(detail=False, methods=['get'])
def inasistencias_frecuentes(self, request):
    '''
    Obtener reporte de afiliados con inasistencias frecuentes.
    URL: GET /api/asistencias/inasistencias_frecuentes/?meses=6&min_inasistencias=3
    '''
    from django.db.models import Count, Q
    from datetime import datetime, timedelta
    from asistencias.models import Asistencia
    
    # Parámetros
    meses = int(request.GET.get('meses', 6))
    min_inasistencias = int(request.GET.get('min_inasistencias', 3))
    
    # Calcular fecha de inicio
    fecha_inicio = datetime.now() - timedelta(days=30 * meses)
    
    # Obtener estadísticas por afiliado
    stats = Asistencia.objects.filter(
        reunion__fecha__gte=fecha_inicio
    ).values(
        'afiliado__id',
        'afiliado__nombres',
        'afiliado__apellidos',
        'afiliado__ci',
        'afiliado__telefono'
    ).annotate(
        total_reuniones=Count('id'),
        total_asistencias=Count('id', filter=Q(presente=True)),
        total_inasistencias=Count('id', filter=Q(presente=False))
    ).filter(
        total_inasistencias__gte=min_inasistencias
    ).order_by('-total_inasistencias')
    
    # Formatear resultados
    resultados = []
    for stat in stats:
        porcentaje_asistencia = (stat['total_asistencias'] / stat['total_reuniones'] * 100) if stat['total_reuniones'] > 0 else 0
        
        # Determinar nivel de alerta
        if porcentaje_asistencia < 50:
            nivel_alerta = 'CRÍTICO'
        elif porcentaje_asistencia < 70:
            nivel_alerta = 'ALTO'
        else:
            nivel_alerta = 'MODERADO'
        
        resultados.append({
            'afiliado_id': stat['afiliado__id'],
            'nombre_completo': f"{stat['afiliado__apellidos']} {stat['afiliado__nombres']}",
            'ci': stat['afiliado__ci'],
            'telefono': stat['afiliado__telefono'],
            'total_reuniones': stat['total_reuniones'],
            'total_asistencias': stat['total_asistencias'],
            'total_inasistencias': stat['total_inasistencias'],
            'porcentaje_asistencia': round(porcentaje_asistencia, 2),
            'nivel_alerta': nivel_alerta
        })
    
    return Response({
        'periodo_meses': meses,
        'fecha_desde': fecha_inicio.date(),
        'total_afiliados_con_problemas': len(resultados),
        'afiliados': resultados
    })
"""

# ========================================
# MÉTODO 3: Estadísticas generales de asistencia
# ========================================
"""
@action(detail=False, methods=['get'])
def estadisticas_generales(self, request):
    '''
    Obtener estadísticas generales de asistencia.
    URL: GET /api/asistencias/estadisticas_generales/
    '''
    from django.db.models import Count, Q, Avg
    from asistencias.models import Asistencia
    
    # Total de asistencias registradas
    total_registros = Asistencia.objects.count()
    total_presentes = Asistencia.objects.filter(presente=True).count()
    total_ausentes = Asistencia.objects.filter(presente=False).count()
    
    porcentaje_asistencia = (total_presentes / total_registros * 100) if total_registros > 0 else 0
    
    # Últimas 5 reuniones
    from reuniones.models import Reunion
    ultimas_reuniones = Reunion.objects.order_by('-fecha')[:5]
    
    reuniones_data = []
    for reunion in ultimas_reuniones:
        asistencias = Asistencia.objects.filter(reunion=reunion)
        total = asistencias.count()
        presentes = asistencias.filter(presente=True).count()
        porcentaje = (presentes / total * 100) if total > 0 else 0
        
        reuniones_data.append({
            'reunion_id': reunion.id,
            'fecha': str(reunion.fecha),
            'tema': getattr(reunion, 'tema', f'Reunión {reunion.id}'),
            'total_convocados': total,
            'total_presentes': presentes,
            'total_ausentes': total - presentes,
            'porcentaje_asistencia': round(porcentaje, 2)
        })
    
    return Response({
        'resumen_general': {
            'total_registros': total_registros,
            'total_presentes': total_presentes,
            'total_ausentes': total_ausentes,
            'porcentaje_asistencia_promedio': round(porcentaje_asistencia, 2)
        },
        'ultimas_reuniones': reuniones_data
    })
"""

# ========================================
# INTEGRACIÓN EN FRONTEND
# ========================================
"""
// En api.js, agregar:
exportarAsistenciasExcel: (params) => {
    const query = new URLSearchParams(params).toString();
    return axios.get(`/asistencias/exportar_excel/?${query}`, { responseType: 'blob' });
},
getInasistenciasFrecuentes: (params) => axios.get('/asistencias/inasistencias_frecuentes/', { params }),
getEstadisticasAsistencia: () => axios.get('/asistencias/estadisticas_generales/'),

// En componente de Asistencias:
const exportarExcel = async () => {
    const response = await api.exportarAsistenciasExcel({
        fecha_desde: '2025-01-01',
        fecha_hasta: '2025-12-31'
    });
    
    const url = window.URL.createObjectURL(new Blob([response.data]));
    const link = document.createElement('a');
    link.href = url;
    link.setAttribute('download', 'asistencias.xlsx');
    document.body.appendChild(link);
    link.click();
    link.remove();
};

// Mostrar inasistencias frecuentes:
const [inasistencias, setInasistencias] = useState([]);

useEffect(() => {
    const loadInasistencias = async () => {
        const res = await api.getInasistenciasFrecuentes({ meses: 6, min_inasistencias: 3 });
        setInasistencias(res.data.afiliados);
    };
    loadInasistencias();
}, []);
"""

print("Código de Asistencia Mejorada listo para integrar")
